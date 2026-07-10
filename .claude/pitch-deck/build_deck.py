#!/usr/bin/env python3
import os, subprocess, html

OUT = "/tmp/claude-0/-home-user-clover/4ef7aa67-e5ff-54aa-b55b-8814ee0bdf78/scratchpad"
CHROME = "/opt/pw-browsers/chromium-1194/chrome-linux/chrome"

CSS = """
:root{--rail:#ede8ff;--stage:#fff;--forest:#1e3d28;--violet:#7c5cbf;--lavender:#d0c8f0;--ink:#3a4a3f;
--serif:'Cormorant Garamond',Georgia,serif;--sans:'DM Sans',system-ui,sans-serif;}
*{box-sizing:border-box;margin:0;padding:0}
html,body{width:1280px;height:720px;overflow:hidden}
body{font-family:var(--sans);color:var(--ink);background:#fff}
.slide{width:1280px;height:720px;background:var(--stage);display:flex;position:relative;container-type:inline-size}
.rail{width:9%;background:var(--rail);display:flex;flex-direction:column;justify-content:space-between;align-items:center;padding:4cqw 0}
.rail .logo{width:4cqw;height:4cqw}
.credit{display:flex;flex-direction:column;align-items:center;text-align:center;line-height:1.25}
.credit .name{font-size:1.3cqw;color:var(--forest);font-weight:500}
.credit .role{font-size:.9cqw;letter-spacing:.22em;color:var(--violet);margin-top:.35cqw}
.stage{flex:1;padding:6cqw 6.5cqw;display:flex;flex-direction:column;justify-content:center}
.stage.has-visual{padding-right:38cqw}
.eyebrow{font-size:1.12cqw;letter-spacing:.2em;text-transform:uppercase;color:var(--violet);font-weight:500;margin-bottom:2.2cqw}
.head{font-family:var(--serif);font-weight:400;color:var(--forest);font-size:4.1cqw;line-height:1.12;max-width:20em}
.head.sm{font-size:3.5cqw}
.brand{font-family:var(--serif);font-weight:400;color:var(--forest);font-size:7.6cqw;line-height:1}
.rule{width:3.6cqw;height:2px;background:var(--lavender);margin:2.6cqw 0}
.italic{font-family:var(--serif);font-style:italic;color:var(--forest);font-size:2.5cqw;line-height:1.3}
.italic.sm{font-size:1.9cqw}
.support{font-size:1.55cqw;color:var(--violet);line-height:1.55;margin-top:1.6cqw;max-width:32em}
svg.ic{stroke:var(--forest);stroke-width:1.4;fill:none;stroke-linecap:round;stroke-linejoin:round}
/* decorative right visuals */
.deco{position:absolute;top:50%;transform:translateY(-50%);display:flex;align-items:center;justify-content:center}
/* chain */
.chain{display:flex;align-items:flex-start;gap:1.1cqw;margin-top:3.6cqw}
.node{display:flex;flex-direction:column;align-items:center;gap:1.1cqw;width:10.5cqw;text-align:center}
.node .iw{width:3.8cqw;height:3.8cqw}
.node .iw svg{width:100%;height:100%}
.node span{font-size:1.2cqw;color:var(--forest);line-height:1.25}
.arrow{color:var(--violet);font-size:2.2cqw;font-family:var(--serif);margin-top:.6cqw}
/* columns */
.cols{display:grid;grid-template-columns:repeat(4,1fr);gap:2.6cqw;margin-top:3.4cqw}
.col{display:flex;flex-direction:column;gap:1.1cqw}
.col .iw{width:4.2cqw;height:4.2cqw}
.col .iw svg{width:100%;height:100%}
.col-label{font-weight:500;color:var(--forest);font-size:1.6cqw}
.col-metric{font-family:var(--serif);font-size:1.5cqw;color:var(--violet);line-height:1.3}
/* steps (product) */
.steps{display:flex;flex-direction:column;gap:2.4cqw;margin-top:3.2cqw}
.step{display:flex;align-items:flex-start;gap:2cqw}
.step .iw{width:3.6cqw;height:3.6cqw;flex:0 0 auto}
.step .iw svg{width:100%;height:100%}
.step .t .l{font-weight:500;color:var(--forest);font-size:1.75cqw}
.step .t .s{font-size:1.35cqw;color:var(--violet);margin-top:.3cqw}
/* closing / ask */
.closing{font-family:var(--serif);font-style:italic;color:var(--forest);font-size:2.9cqw;line-height:1.28;margin-top:3cqw;max-width:20em}
.ask{display:flex;gap:2.4cqw;margin-top:2.6cqw;font-size:1.55cqw;color:var(--violet);letter-spacing:.03em}
.ask span{display:flex;align-items:center;gap:.7cqw}
.ask .dot{width:.55cqw;height:.55cqw;border-radius:50%;background:var(--violet)}
/* image slot */
.slot{position:absolute;border:2px dashed var(--lavender);border-radius:14px;background:#f5f2fb;
display:flex;flex-direction:column;align-items:center;justify-content:center;gap:1.2cqw}
.slot svg{width:3.4cqw;height:3.4cqw;stroke:var(--violet);stroke-width:1.3;fill:none;stroke-linecap:round;stroke-linejoin:round}
.slot .sl{font-size:1.35cqw;color:var(--violet);letter-spacing:.03em;font-weight:500}
.slot .ss{font-size:1.05cqw;color:#a89bd0}
"""

def logo(center="#ede8ff"):
    return ('<svg class="logo" viewBox="0 0 40 40"><g fill="#1e3d28">'
            '<circle cx="20" cy="12" r="8"/><circle cx="20" cy="28" r="8"/>'
            '<circle cx="12" cy="20" r="8"/><circle cx="28" cy="20" r="8"/></g>'
            f'<circle cx="20" cy="20" r="5" fill="{center}"/></svg>')

RAIL = ('<aside class="rail">' + logo() +
        '<div class="credit"><span class="name">Chloe Guo</span>'
        '<span class="role">FOUNDER</span></div></aside>')

# icon library (inner svg paths)
IC = {
 "user":'<path d="M20 21v-2a4 4 0 0 0-4-4H8a4 4 0 0 0-4 4v2"/><circle cx="12" cy="7" r="4"/>',
 "link":'<path d="M10 13a5 5 0 0 0 7.54.54l3-3a5 5 0 0 0-7.07-7.07l-1.72 1.71"/><path d="M14 11a5 5 0 0 0-7.54-.54l-3 3a5 5 0 0 0 7.07 7.07l1.71-1.71"/>',
 "check":'<path d="M22 11.08V12a10 10 0 1 1-5.93-9.14"/><polyline points="22 4 12 14.01 9 11.01"/>',
 "zap":'<polygon points="13 2 3 14 12 14 11 22 21 10 12 10 13 2"/>',
 "building":'<path d="M3 21h18"/><path d="M5 21V7l7-4 7 4v14"/><path d="M9.5 21v-5h5v5"/>',
 "globe":'<circle cx="12" cy="12" r="10"/><line x1="2" y1="12" x2="22" y2="12"/><path d="M12 2a15.3 15.3 0 0 1 4 10 15.3 15.3 0 0 1-4 10 15.3 15.3 0 0 1-4-10 15.3 15.3 0 0 1 4-10z"/>',
 "compass":'<circle cx="12" cy="12" r="10"/><polygon points="16.24 7.76 14.12 14.12 7.76 16.24 9.88 9.88 16.24 7.76"/>',
 "eye":'<path d="M1 12s4-8 11-8 11 8 11 8-4 8-11 8-11-8-11-8z"/><circle cx="12" cy="12" r="3"/>',
 "coffee":'<path d="M18 8h1a4 4 0 0 1 0 8h-1"/><path d="M2 8h16v9a4 4 0 0 1-4 4H6a4 4 0 0 1-4-4V8z"/><line x1="6" y1="2" x2="6" y2="5"/><line x1="10" y1="2" x2="10" y2="5"/><line x1="14" y1="2" x2="14" y2="5"/>',
 "cube":'<path d="M21 16V8a2 2 0 0 0-1-1.73l-7-4a2 2 0 0 0-2 0l-7 4A2 2 0 0 0 3 8v8a2 2 0 0 0 1 1.73l7 4a2 2 0 0 0 2 0l7-4A2 2 0 0 0 21 16z"/><polyline points="3.27 6.96 12 12.01 20.73 6.96"/><line x1="12" y1="22.08" x2="12" y2="12"/>',
 "msg":'<path d="M21 15a2 2 0 0 1-2 2H7l-4 4V5a2 2 0 0 1 2-2h14a2 2 0 0 1 2 2z"/>',
 "trend":'<polyline points="23 6 13.5 15.5 8.5 10.5 1 18"/><polyline points="17 6 23 6 23 12"/>',
 "img":'<rect x="3" y="3" width="18" height="18" rx="2"/><circle cx="8.5" cy="8.5" r="1.5"/><polyline points="21 15 16 10 5 21"/>',
}
def ic(name, cls="ic"):
    return f'<svg class="{cls}" viewBox="0 0 24 24">{IC[name]}</svg>'

# image slots rendered as SEPARATE replaceable pictures: idx -> (x,y,w,h,label,sub)
SLOTS = {
 3:(795,150,385,410,"Café photo","Two founders, three feet apart"),
 5:(905,115,250,490,"Product screenshot","Discover · Learn · Invite"),
 7:(835,130,330,450,"Founder portrait","Chloe Guo"),
}

# ---- slides: (stage_html, slot_html) ----
slides = []

# 1 cover
slides.append((
 f'<div class="brand">Clover</div><div class="rule"></div>'
 f'<p class="italic">We’ve never been more connected, or more alone.</p>'
 f'<p class="support">Your people are already here.<br>Clover is how you find them.</p>',
 '<div class="deco" style="right:9cqw"><svg viewBox="0 0 40 40" style="width:20cqw;height:20cqw">'
 '<g fill="#1e3d28"><circle cx="20" cy="12" r="8"/><circle cx="20" cy="28" r="8"/>'
 '<circle cx="12" cy="20" r="8"/><circle cx="28" cy="20" r="8"/></g>'
 '<circle cx="20" cy="20" r="5" fill="#ffffff"/></svg></div>'))

# 2 paradox
slides.append((
 '<div class="eyebrow">01 — The Paradox</div>'
 '<h2 class="head">More people are building than ever. Almost all of them are building alone.</h2>'
 '<div class="rule"></div>'
 '<p class="italic sm">The tools got smarter. The teams got smaller. But building became lonelier.</p>',
 '<div class="deco" style="right:11cqw"><svg class="ic" viewBox="0 0 48 48" style="width:15cqw;height:15cqw">'
 '<circle cx="24" cy="24" r="22" stroke="#d0c8f0" stroke-width="1"/>'
 '<path d="M31 33v-1.5a4.5 4.5 0 0 0-4.5-4.5h-5A4.5 4.5 0 0 0 17 31.5V33"/>'
 '<circle cx="24" cy="19" r="4.5"/></svg></div>'))

# 3 moment + cafe slot
slides.append((
 '<div class="eyebrow">02 — The Moment</div>'
 '<h2 class="head">The next great company might start three feet away — and never happen.</h2>'
 '<div class="rule"></div>'
 '<p class="italic sm">Two founders. Same café. Same problem.</p>'
 '<p class="support">They leave without ever knowing what could have happened.</p>',
 ''))

# 4 chain
slides.append((
 '<div class="eyebrow">03 — Why It Matters</div>'
 '<h2 class="head sm">The strongest startup ecosystems are built one relationship at a time.</h2>'
 '<div class="chain">'
 f'<div class="node"><div class="iw">{ic("link")}</div><span>Meaningful<br>connection</span></div>'
 '<div class="arrow">→</div>'
 f'<div class="node"><div class="iw">{ic("check")}</div><span>Accountability</span></div>'
 '<div class="arrow">→</div>'
 f'<div class="node"><div class="iw">{ic("zap")}</div><span>Better<br>execution</span></div>'
 '<div class="arrow">→</div>'
 f'<div class="node"><div class="iw">{ic("building")}</div><span>Better<br>companies</span></div>'
 '<div class="arrow">→</div>'
 f'<div class="node"><div class="iw">{ic("globe")}</div><span>Stronger innovation ecosystems</span></div>'
 '</div>',
 ''))

# 5 product + steps + phone slot
slides.append((
 '<div class="eyebrow">04 — The Product</div>'
 '<h2 class="head sm">Clover turns chance encounters into meaningful founder relationships.</h2>'
 '<div class="rule"></div>'
 '<div class="steps">'
 f'<div class="step"><div class="iw">{ic("compass")}</div><div class="t"><div class="l">Discover</div><div class="s">Who’s building near you, right now.</div></div></div>'
 f'<div class="step"><div class="iw">{ic("eye")}</div><div class="t"><div class="l">Learn</div><div class="s">What they’re working on, before hello.</div></div></div>'
 f'<div class="step"><div class="iw">{ic("coffee")}</div><div class="t"><div class="l">Invite</div><div class="s">Cowork together — today.</div></div></div>'
 '</div>',
 ''))

# 6 execution 4 cols
slides.append((
 '<div class="eyebrow">05 — Execution</div>'
 '<h2 class="head">Execution before scale.</h2>'
 '<div class="rule"></div>'
 '<div class="cols">'
 f'<div class="col"><div class="iw">{ic("cube")}</div><div class="col-label">Build</div><div class="col-metric">Solo, in four months</div></div>'
 f'<div class="col"><div class="iw">{ic("msg")}</div><div class="col-label">Learn</div><div class="col-metric">Every founder interviewed</div></div>'
 f'<div class="col"><div class="iw">{ic("trend")}</div><div class="col-label">Grow</div><div class="col-metric">26 users, all organic</div></div>'
 f'<div class="col"><div class="iw">{ic("link")}</div><div class="col-label">Connect</div><div class="col-metric">Matched by hand</div></div>'
 '</div>',
 ''))

# 7 why me + portrait slot
slides.append((
 '<div class="eyebrow">06 — Why Me</div>'
 '<h2 class="head">I’m building the bridge I’ve lived on my whole life.</h2>'
 '<div class="rule"></div>'
 '<p class="italic sm">Raised in China. Built in Canada.</p>'
 '<p class="support">I move between the worlds Clover wants to connect.</p>',
 ''))

# 8 vision & ask
slides.append((
 '<div class="eyebrow">07 — Vision &amp; Ask</div>'
 '<h2 class="head sm">The next great company starts with two builders in the same room. Clover puts them there — on purpose.</h2>'
 '<p class="closing">Great ideas can come from anywhere.<br>Great companies are built together.</p>'
 '<div class="ask"><span><span class="dot"></span>Users</span>'
 '<span><span class="dot"></span>Advisors</span>'
 '<span><span class="dot"></span>Partners</span></div>',
 ''))

FONTS = ('<link href="https://fonts.googleapis.com/css2?family=Cormorant+Garamond:'
         'ital,wght@0,300;0,400;0,500;0,600;1,300;1,400;1,500&family=DM+Sans:wght@300;400;500&display=swap" rel="stylesheet">')
PAGE = ('<!doctype html><html><head><meta charset="utf-8">'+FONTS+
        '<style>%s</style></head><body>'
        '<div class="slide">%s<main class="stage %s">%s</main></div></body></html>')

def render(html_str, w, h, outpng):
    hp = outpng.replace(".png",".html")
    open(hp,"w").write(html_str)
    subprocess.run([CHROME,"--headless","--no-sandbox","--disable-gpu","--hide-scrollbars",
        "--force-device-scale-factor=2", f"--window-size={w},{h}",
        "--virtual-time-budget=2500","--default-background-color=FFFFFFFF",
        f"--screenshot={outpng}", f"file://{hp}"],
        stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)

SLOT_CSS = ("*{margin:0;box-sizing:border-box}html,body{width:%dpx;height:%dpx}"
 "body{font-family:'DM Sans',system-ui,sans-serif}"
 ".slot{width:100%%;height:100%%;border:2px dashed #d0c8f0;border-radius:14px;background:#f5f2fb;"
 "display:flex;flex-direction:column;align-items:center;justify-content:center;gap:14px}"
 ".slot svg{width:44px;height:44px;stroke:#7c5cbf;stroke-width:1.3;fill:none;stroke-linecap:round;stroke-linejoin:round}"
 ".slot .sl{font-size:17px;color:#7c5cbf;font-weight:500;letter-spacing:.02em}"
 ".slot .ss{font-size:13px;color:#a89bd0}")

pngs=[]
for i,(stage,_) in enumerate(slides,1):
    has_visual = "has-visual" if i in SLOTS else ""
    doc = PAGE % (CSS, RAIL, has_visual, stage)
    pp = os.path.join(OUT, f"slide{i}.png")
    render(doc, 1280, 720, pp)
    pngs.append(pp)
    print("rendered", pp, os.path.getsize(pp))

# render replaceable slot placeholder pictures
slot_pngs={}
for i,(x,y,w,h,label,sub) in SLOTS.items():
    doc=('<!doctype html><html><head><meta charset="utf-8">'+FONTS+
         '<style>'+(SLOT_CSS%(w,h))+'</style></head><body><div class="slot">'
         f'<svg viewBox="0 0 24 24">{IC["img"]}</svg>'
         f'<div class="sl">{label}</div><div class="ss">{sub}</div></div></body></html>')
    pp=os.path.join(OUT,f"slot{i}.png")
    render(doc,w,h,pp)
    slot_pngs[i]=pp
    print("slot",pp,os.path.getsize(pp))

# ---- assemble pptx ----
from pptx import Presentation
from pptx.util import Emu
EMU=9525  # per css px at 1280-wide / 13.333in slide
prs = Presentation()
prs.slide_width = Emu(12192000)
prs.slide_height = Emu(6858000)
blank = prs.slide_layouts[6]
for idx,pp in enumerate(pngs,1):
    s = prs.slides.add_slide(blank)
    s.shapes.add_picture(pp, 0, 0, width=prs.slide_width, height=prs.slide_height)
    if idx in SLOTS:
        x,y,w,h,_,_ = SLOTS[idx]
        pic = s.shapes.add_picture(slot_pngs[idx], Emu(x*EMU), Emu(y*EMU),
                                   width=Emu(w*EMU), height=Emu(h*EMU))
        pic.name = SLOTS[idx][4]  # e.g. "Café photo"
out = os.path.join(OUT, "Clover-Pitch-Deck.pptx")
prs.save(out)
print("SAVED", out, os.path.getsize(out))
